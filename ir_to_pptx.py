"""GraphIR → 可编辑 PPTX（python-pptx 原生形状）。缺失依赖不阻断四件套。

T14.8：不要直接套 1180px 坐标。按 16:9 画布等比缩放并居中，尺寸来自 design_tokens.page。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

try:
    from design_tokens import page_config, shape_colors
except Exception:  # noqa: BLE001
    def page_config() -> dict[str, Any]:
        return {"pptx_width_in": 13.333, "pptx_height_in": 7.5, "pptx_margin_in": 0.4}

    def shape_colors() -> dict[str, str]:
        return {"fill": "#e7effe", "stroke": "#6ea8dc", "text": "#17427e"}


def _hex_rgb(value: str) -> tuple[int, int, int]:
    raw = (value or "").strip().lstrip("#")
    if len(raw) == 3:
        raw = "".join(ch * 2 for ch in raw)
    if len(raw) != 6:
        raw = "17427e"
    return int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)


def _canvas(ir: dict[str, Any]) -> dict[str, float]:
    page = page_config()
    slide_w = float(page.get("pptx_width_in") or 13.333)
    slide_h = float(page.get("pptx_height_in") or 7.5)
    margin = float(page.get("pptx_margin_in") or 0.4)
    size = ir.get("size") or {}
    ir_w = max(float(size.get("w") or 1), 1.0)
    ir_h = max(float(size.get("h") or 1), 1.0)
    # CSS px → inch（96dpi），再按 16:9 可用区域等比缩放并居中
    ir_w_in = ir_w / 96.0
    ir_h_in = ir_h / 96.0
    avail_w = max(slide_w - margin * 2, 0.5)
    avail_h = max(slide_h - margin * 2, 0.5)
    scale = min(avail_w / ir_w_in, avail_h / ir_h_in)
    used_w = ir_w_in * scale
    used_h = ir_h_in * scale
    return {
        "slide_w": slide_w,
        "slide_h": slide_h,
        "scale": scale,
        "ox": margin + (avail_w - used_w) / 2,
        "oy": margin + (avail_h - used_h) / 2,
    }


def map_box(node: dict[str, Any], canvas: dict[str, float]) -> tuple[float, float, float, float]:
    scale = canvas["scale"] / 96.0
    return (
        canvas["ox"] + float(node.get("x") or 0) * scale,
        canvas["oy"] + float(node.get("y") or 0) * scale,
        max(float(node.get("w") or 20) * scale, 0.15),
        max(float(node.get("h") or 12) * scale, 0.12),
    )


def write_pptx(ir: dict[str, Any], path: Path) -> Path:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError(
            "缺少 python-pptx，无法导出可编辑 PPT。"
            "请执行：python3 -m pip install python-pptx"
            "（四件套 png/html/md/mmd 不受影响）"
        ) from exc

    canvas = _canvas(ir)
    colors = shape_colors()
    prs = Presentation()
    prs.slide_width = Inches(canvas["slide_w"])
    prs.slide_height = Inches(canvas["slide_h"])
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    title = str(ir.get("title") or "").strip()
    if title:
        box = slide.shapes.add_textbox(
            Inches(0.4), Inches(0.12), Inches(canvas["slide_w"] - 0.8), Inches(0.36)
        )
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*_hex_rgb(colors.get("text") or "#17427e"))

    id_to_center: dict[str, tuple[float, float]] = {}
    for node in ir.get("nodes") or []:
        x, y, w, h = map_box(node, canvas)
        kind = str(node.get("shape") or "round")
        mso = MSO_SHAPE.OVAL if kind == "ellipse" else MSO_SHAPE.ROUNDED_RECTANGLE
        if kind == "bar":
            mso = MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(mso, Inches(x), Inches(y), Inches(w), Inches(h))
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*_hex_rgb(str(node.get("fill") or colors.get("fill"))))
        shape.line.color.rgb = RGBColor(*_hex_rgb(str(node.get("stroke") or colors.get("stroke"))))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        try:
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            shape.text_frame.paragraphs[0].space_before = Pt(0)
            tf.anchor = MSO_ANCHOR.MIDDLE
        except Exception:  # noqa: BLE001
            pass
        text = str(node.get("text") or "")
        tf.clear()
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        font_px = float(node.get("font_px") or 12)
        p.font.size = Pt(max(8, font_px * 0.75 * canvas["scale"]))
        p.font.color.rgb = RGBColor(*_hex_rgb(colors.get("text") or "#17427e"))
        nid = str(node.get("id") or "")
        if nid:
            id_to_center[nid] = (x + w / 2, y + h / 2)

    for edge in ir.get("edges") or []:
        src = id_to_center.get(str(edge.get("from")))
        dst = id_to_center.get(str(edge.get("to")))
        if not src or not dst:
            continue
        try:
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(src[0]),
                Inches(src[1]),
                Inches(dst[0]),
                Inches(dst[1]),
            )
            conn.line.color.rgb = RGBColor(*_hex_rgb(colors.get("stroke") or "#6ea8dc"))
            if edge.get("dashed"):
                try:
                    from pptx.enum.dml import MSO_LINE

                    conn.line.dash_style = MSO_LINE.DASH
                except Exception:  # noqa: BLE001
                    pass
        except Exception:  # noqa: BLE001
            continue

    path = Path(path)
    prs.save(str(path))
    return path
